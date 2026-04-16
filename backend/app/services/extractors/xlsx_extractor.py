"""
Extract translatable cell text from .xlsx files.

Rules (v1):
- Translate only cells whose value is a non-empty string and not a formula.
- Keep formulas as-is.
- Preserve sheet structure; rebuild sets cell values back in the same positions.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.services.extractors.base import ExtractionResult, Segment, SegmentType


def extract_xlsx(path: Path) -> ExtractionResult:
    wb = load_workbook(path)

    segments: list[Segment] = []
    cells_meta: list[dict] = []
    idx = 0

    for ws in wb.worksheets:
        sheet_name = ws.title
        # iter_rows with values_only=False gives Cell objects
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val is None:
                    continue
                if isinstance(val, str):
                    # openpyxl represents formulas with data_type == "f" and/or leading '='
                    if cell.data_type == "f" or val.strip().startswith("="):
                        continue
                    if not val.strip():
                        continue

                    segments.append(
                        Segment(
                            index=idx,
                            content=val,
                            segment_type=SegmentType.CELL,
                            meta={},
                        )
                    )
                    cells_meta.append(
                        {"sheet": sheet_name, "row": cell.row, "col": cell.column}
                    )
                    idx += 1

    return ExtractionResult(
        segments=segments,
        format="xlsx",
        meta={"cells": cells_meta},
    )

