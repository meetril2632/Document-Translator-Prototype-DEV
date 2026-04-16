"""
Extract text from DOCX files using python-docx.

Preserves paragraph and run structure so we can rebuild with same formatting.
"""

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from app.services.extractors.base import ExtractionResult, Segment, SegmentType


def _iter_block_items(parent):
    """
    Yield each paragraph and table child within parent (document or cell).
    From python-docx docs pattern for iterating in order.
    """
    if isinstance(parent, DocumentType):
        for child in parent.element.body:
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif child.tag.endswith("tbl"):
                yield Table(child, parent)
    else:
        for child in parent._element:
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif child.tag.endswith("tbl"):
                yield Table(child, parent)


def _iter_docx_blocks(doc: DocumentType):
    """
    Yield paragraphs and tables in a stable order including:
    - all section headers (paragraphs, then tables)
    - main document body (paragraphs/tables interleaved as in XML)
    - all section footers (paragraphs, then tables)

    The rebuilder uses the same traversal so segment-to-text mapping stays correct.
    """
    # Headers
    for section in doc.sections:
        header = section.header
        for p in header.paragraphs:
            yield p
        for t in header.tables:
            yield t

    # Body (XML interleaving)
    for block in _iter_block_items(doc):
        yield block

    # Footers
    for section in doc.sections:
        footer = section.footer
        for p in footer.paragraphs:
            yield p
        for t in footer.tables:
            yield t


def _paragraph_to_segments(paragraph: Paragraph, start_index: int) -> list[Segment]:
    """
    Split a paragraph into runs (text with same style). Each run becomes one segment.
    """
    segments = []
    for run in paragraph.runs:
        if run.text.strip():
            segments.append(
                Segment(
                    index=start_index,
                    content=run.text,
                    segment_type=SegmentType.TEXT,
                    meta={"style": run.style.name if run.style else None},
                )
            )
            start_index += 1
    # If paragraph had no runs or only empty runs, still add one segment for the full text
    if not segments and paragraph.text.strip():
        segments.append(
            Segment(
                index=start_index,
                content=paragraph.text,
                segment_type=SegmentType.TEXT,
                meta={"style": paragraph.style.name if paragraph.style else None},
            )
        )
    return segments


def extract_docx(path: Path) -> ExtractionResult:
    """
    Extract all text from DOCX in order: paragraphs and table cells.
    Each run (or paragraph if no runs) is one segment for translation.

    :param path: Path to the .docx file
    :return: ExtractionResult with segments of type TEXT and format "docx"
    """
    doc = Document(path)
    segments = []
    idx = 0

    for block in _iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            for seg in _paragraph_to_segments(block, idx):
                seg.index = idx
                segments.append(seg)
                idx += 1
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for seg in _paragraph_to_segments(para, idx):
                            seg.index = idx
                            segments.append(seg)
                            idx += 1

    return ExtractionResult(
        segments=segments,
        format="docx",
        meta={},  # Rebuilder will reopen file and replace text in order
    )
