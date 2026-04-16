"""
Rebuild DOCX by replacing text in runs/paragraphs with translated content.

Reopens the source document, walks the same structure as the extractor,
and replaces each segment's text with the corresponding translated string.
"""

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.extractors.docx_extractor import _iter_docx_blocks


def rebuild_docx(
    source_path: Path,
    output_path: Path,
    translated_contents: list[str],
    meta: dict,
) -> None:
    """
    Create a new DOCX with the same structure as source_path but with
    paragraph/run text replaced by translated_contents (same order as extraction).

    :param source_path: Original DOCX path
    :param output_path: Where to write the translated DOCX
    :param translated_contents: List of translated strings in extraction order
    :param meta: Unused for docx (structure comes from source)
    """
    doc = Document(source_path)
    idx = [0]  # Use list so we can mutate inside nested functions

    def set_paragraph_text(paragraph: Paragraph):
        """Replace paragraph content with next translated string(s). One translation per segment (run or full para)."""
        runs = list(paragraph.runs)
        if runs:
            for run in runs:
                # Only consume a translation for runs that had text (same as extractor)
                if run.text.strip() and idx[0] < len(translated_contents):
                    run.text = translated_contents[idx[0]]
                    idx[0] += 1
        else:
            if paragraph.text.strip() and idx[0] < len(translated_contents):
                paragraph.clear()
                paragraph.add_run(translated_contents[idx[0]])
                idx[0] += 1

    for block in _iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            set_paragraph_text(block)
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        set_paragraph_text(para)

    doc.save(output_path)
