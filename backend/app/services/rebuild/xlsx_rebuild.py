"""
Rebuild .xlsx by replacing cell string values with translated content.

This preserves:
- sheet structure
- cell styles (we only set .value)
- formulas (we do not touch formula cells because extraction skipped them)
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


def rebuild_xlsx(
    source_path: Path,
    output_path: Path,
    translated_contents: list[str],
    meta: dict,
) -> None:
    wb = load_workbook(source_path)
    cells_meta: list[dict] = meta.get("cells", [])

    if len(cells_meta) != len(translated_contents):
        raise ValueError(
            f"XLSX rebuild mismatch: {len(cells_meta)} cells meta vs {len(translated_contents)} translated strings"
        )

    for i, pos in enumerate(cells_meta):
        sheet_name = pos["sheet"]
        row = int(pos["row"])
        col = int(pos["col"])
        ws = wb[sheet_name]
        ws.cell(row=row, column=col).value = translated_contents[i]

    wb.save(output_path)

